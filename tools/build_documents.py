#!/usr/bin/env python3
"""Build DOCX and PDF versions of CyberSafe_report.md without external converters."""

from __future__ import annotations

from dataclasses import dataclass
from html import escape
from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image,
    KeepTogether,
    LongTable,
    PageBreak,
    Paragraph,
    Preformatted,
    SimpleDocTemplate,
    Spacer,
    TableStyle,
)

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "CyberSafe_report.md"
DOCX = ROOT / "CyberSafe_LCT2026_jury_final.docx"
PDF = ROOT / "CyberSafe_LCT2026_jury_final.pdf"
PURPLE = "4C1D95"
DARK = "1F2937"
LIGHT_PURPLE = "EDE9FE"
LIGHT_GRAY = "F3F4F6"


@dataclass
class Block:
    kind: str
    value: object
    level: int = 0


def parse_markdown(text: str) -> list[Block]:
    lines = text.splitlines()
    blocks: list[Block] = []
    i = 0
    paragraph: list[str] = []

    def flush() -> None:
        nonlocal paragraph
        if paragraph:
            blocks.append(Block("paragraph", " ".join(x.strip() for x in paragraph)))
            paragraph = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped == "<!-- PAGEBREAK -->":
            flush()
            blocks.append(Block("pagebreak", ""))
            i += 1
            continue
        if stripped.startswith("```"):
            flush()
            language = stripped[3:].strip()
            i += 1
            code: list[str] = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            if i < len(lines):
                i += 1
            blocks.append(Block("code", "\n".join(code), 0 if not language else 1))
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            flush()
            blocks.append(Block("heading", m.group(2).strip(), len(m.group(1))))
            i += 1
            continue
        m = re.match(r"^!\[(.*?)\]\((.*?)\)\s*$", stripped)
        if m:
            flush()
            blocks.append(Block("image", (m.group(1), m.group(2))))
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{3,}", lines[i + 1]):
            flush()
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                rows.append(row)
                i += 1
            if len(rows) >= 2:
                rows.pop(1)
            blocks.append(Block("table", rows))
            continue
        m = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        if m:
            flush()
            blocks.append(Block("bullet", m.group(2), len(m.group(1)) // 2))
            i += 1
            continue
        m = re.match(r"^(\s*)(\d+)\.\s+(.*)$", line)
        if m:
            flush()
            blocks.append(Block("number", (m.group(2), m.group(3)), len(m.group(1)) // 2))
            i += 1
            continue
        if stripped == "":
            flush()
            i += 1
            continue
        paragraph.append(line)
        i += 1
    flush()
    return blocks


LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
TOKEN_RE = re.compile(r"(\*\*.*?\*\*|`.*?`|\[[^\]]+\]\([^)]+\))")


def plain_inline(text: str) -> str:
    text = LINK_RE.sub(lambda m: f"{m.group(1)} ({m.group(2)})", text)
    return text.replace("**", "").replace("`", "")


def add_docx_runs(paragraph, text: str) -> None:
    pos = 0
    for match in TOKEN_RE.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(76, 29, 149)
        else:
            lm = LINK_RE.fullmatch(token)
            paragraph.add_run(f"{lm.group(1)} ({lm.group(2)})")
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Страница ")
    run.font.size = Pt(8)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


def set_docx_cell_text(cell, text: str, bold: bool = False, size: float = 8.0) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(plain_inline(text))
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def build_docx(blocks: list[Block]) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.6)
    sec.bottom_margin = Cm(1.6)
    sec.left_margin = Cm(1.7)
    sec.right_margin = Cm(1.7)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.06
    for name, size, color in [
        ("Title", 24, PURPLE),
        ("Heading 1", 17, PURPLE),
        ("Heading 2", 14, PURPLE),
        ("Heading 3", 11.5, DARK),
    ]:
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    header = sec.header.paragraphs[0]
    header.text = "CyberSafe — отчёт по реверсу"
    header.runs[0].font.name = "Arial"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor(107, 114, 128)
    add_page_field(sec.footer.paragraphs[0])

    first_h1 = True
    for block in blocks:
        if block.kind == "pagebreak":
            doc.add_page_break()
        elif block.kind == "heading":
            level = block.level
            if level == 1 and first_h1:
                p = doc.add_paragraph(style="Title")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_docx_runs(p, str(block.value))
                p.paragraph_format.space_before = Pt(30)
                p.paragraph_format.space_after = Pt(18)
                first_h1 = False
            else:
                style = "Heading 1" if level == 2 else "Heading 2" if level == 3 else "Heading 3"
                p = doc.add_paragraph(style=style)
                add_docx_runs(p, str(block.value))
        elif block.kind == "paragraph":
            p = doc.add_paragraph()
            add_docx_runs(p, str(block.value))
        elif block.kind in ("bullet", "number"):
            text = block.value if block.kind == "bullet" else block.value[1]
            style = "List Bullet" if block.kind == "bullet" else "List Number"
            p = doc.add_paragraph(style=style)
            p.paragraph_format.left_indent = Cm(0.65 + 0.45 * block.level)
            p.paragraph_format.first_line_indent = Cm(-0.3)
            add_docx_runs(p, str(text))
        elif block.kind == "code":
            table = doc.add_table(rows=1, cols=1)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = table.cell(0, 0)
            shade_cell(cell, LIGHT_GRAY)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(block.value))
            run.font.name = "Courier New"
            run.font.size = Pt(7.7)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
        elif block.kind == "image":
            alt, rel = block.value
            path = ROOT / rel
            if path.exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run()
                run.add_picture(str(path), width=Cm(15.8))
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cr = cap.add_run(alt)
                cr.italic = True
                cr.font.name = "Arial"
                cr.font.size = Pt(8)
                cr.font.color.rgb = RGBColor(107, 114, 128)
        elif block.kind == "table":
            rows = block.value
            if not rows:
                continue
            cols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.style = "Table Grid"
            table.autofit = True
            repeat_table_header(table.rows[0])
            for ri, row in enumerate(rows):
                for ci in range(cols):
                    txt = row[ci] if ci < len(row) else ""
                    set_docx_cell_text(table.cell(ri, ci), txt, bold=ri == 0, size=7.0 if cols >= 5 else 8.2)
                    if ri == 0:
                        shade_cell(table.cell(ri, ci), LIGHT_PURPLE)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)

    props = doc.core_properties
    props.title = "CyberSafe — отчёт по реверсу и получению доступа"
    props.subject = "Лидеры цифровой трансформации 2026, кейс Positive Technologies"
    props.author = "Команда участника"
    props.keywords = "CyberSafe, RP2040, PICOBOOT, reverse engineering, LCT"
    doc.save(DOCX)


def pdf_inline(text: str) -> str:
    placeholders: list[str] = []

    def stash(value: str) -> str:
        placeholders.append(value)
        return f"@@TOKEN{len(placeholders)-1}@@"

    text = LINK_RE.sub(lambda m: stash(f"{m.group(1)} ({m.group(2)})"), text)
    text = re.sub(r"`([^`]+)`", lambda m: stash(f'<font name="DejaVuMono" color="#4c1d95">{escape(m.group(1))}</font>'), text)
    text = escape(text)
    text = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)
    for idx, token in enumerate(placeholders):
        replacement = token if token.startswith("<font") else escape(token)
        text = text.replace(f"@@TOKEN{idx}@@", replacement)
    return text


def table_widths(rows: list[list[str]], available: float) -> list[float]:
    cols = max(len(r) for r in rows)
    if cols == 6:
        weights = [0.45, 3.0, 1.35, 0.9, 0.55, 1.9]
    elif cols == 5:
        weights = [1.2, 0.85, 1.35, 1.35, 2.0]
    elif cols == 3:
        weights = [2.0, 1.0, 3.4]
    elif cols == 2:
        weights = [2.2, 4.2]
    else:
        weights = [1.0] * cols
    scale = available / sum(weights)
    return [w * scale for w in weights]


def build_pdf(blocks: list[Block]) -> None:
    regular = "/System/Library/Fonts/Supplemental/Arial.ttf"
    bold = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
    italic = "/System/Library/Fonts/Supplemental/Arial Italic.ttf"
    mono = "/System/Library/Fonts/Supplemental/Courier New.ttf"
    pdfmetrics.registerFont(TTFont("ArialReport", regular))
    pdfmetrics.registerFont(TTFont("ArialReportBold", bold))
    if Path(italic).exists():
        pdfmetrics.registerFont(TTFont("ArialReportItalic", italic))
    else:
        pdfmetrics.registerFont(TTFont("ArialReportItalic", regular))
    pdfmetrics.registerFont(TTFont("DejaVuMono", mono))
    pdfmetrics.registerFontFamily("ArialReport", normal="ArialReport", bold="ArialReportBold", italic="ArialReportItalic")

    doc = SimpleDocTemplate(
        str(PDF),
        pagesize=A4,
        rightMargin=1.35 * cm,
        leftMargin=1.35 * cm,
        topMargin=1.45 * cm,
        bottomMargin=1.45 * cm,
        title="CyberSafe — отчёт по реверсу",
        author="Команда участника",
    )
    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "BodyRU", parent=styles["BodyText"], fontName="ArialReport", fontSize=9.2,
        leading=11.4, textColor=colors.HexColor("#1f2937"), spaceAfter=4,
    )
    title = ParagraphStyle(
        "TitleRU", parent=body, fontName="ArialReportBold", fontSize=21,
        leading=25, alignment=TA_CENTER, textColor=colors.HexColor("#4c1d95"),
        spaceBefore=25, spaceAfter=18,
    )
    h1 = ParagraphStyle(
        "H1RU", parent=body, fontName="ArialReportBold", fontSize=15,
        leading=18, textColor=colors.HexColor("#4c1d95"), spaceBefore=10, spaceAfter=6,
        keepWithNext=True,
    )
    h1_before_table = ParagraphStyle("H1BeforeTableRU", parent=h1, keepWithNext=False)
    h2 = ParagraphStyle(
        "H2RU", parent=body, fontName="ArialReportBold", fontSize=12,
        leading=15, textColor=colors.HexColor("#4c1d95"), spaceBefore=8, spaceAfter=5,
        keepWithNext=True,
    )
    h3 = ParagraphStyle(
        "H3RU", parent=body, fontName="ArialReportBold", fontSize=10.3,
        leading=13, textColor=colors.HexColor("#374151"), spaceBefore=7, spaceAfter=4,
        keepWithNext=True,
    )
    bullet = ParagraphStyle(
        "BulletRU", parent=body, leftIndent=13, firstLineIndent=-8, bulletIndent=3, spaceAfter=2.5,
    )
    code = ParagraphStyle(
        "CodeRU", parent=body, fontName="DejaVuMono", fontSize=6.9, leading=8.5,
        leftIndent=6, rightIndent=6, borderPadding=5, backColor=colors.HexColor("#f3f4f6"),
        spaceBefore=3, spaceAfter=6,
    )
    caption = ParagraphStyle(
        "CaptionRU", parent=body, fontName="ArialReportItalic", fontSize=7.8,
        leading=9, alignment=TA_CENTER, textColor=colors.HexColor("#6b7280"), spaceAfter=7,
    )
    table_head = ParagraphStyle(
        "TableHeadRU", parent=body, fontName="ArialReportBold", fontSize=7.4, leading=8.7,
        textColor=colors.HexColor("#1f2937"),
    )
    table_cell = ParagraphStyle(
        "TableCellRU", parent=body, fontName="ArialReport", fontSize=7.2, leading=8.5,
        textColor=colors.HexColor("#1f2937"),
    )

    story = []
    first_h1 = True
    for block_index, block in enumerate(blocks):
        if block.kind == "pagebreak":
            story.append(PageBreak())
        elif block.kind == "heading":
            if block.level == 1 and first_h1:
                story.append(Paragraph(pdf_inline(str(block.value)), title))
                first_h1 = False
            else:
                next_is_table = block_index + 1 < len(blocks) and blocks[block_index + 1].kind == "table"
                if block.level == 2:
                    style = h1_before_table if next_is_table else h1
                else:
                    style = h2 if block.level == 3 else h3
                story.append(Paragraph(pdf_inline(str(block.value)), style))
        elif block.kind == "paragraph":
            story.append(Paragraph(pdf_inline(str(block.value)), body))
        elif block.kind == "bullet":
            story.append(Paragraph("• " + pdf_inline(str(block.value)), bullet))
        elif block.kind == "number":
            number, text = block.value
            story.append(Paragraph(f"{number}. " + pdf_inline(str(text)), bullet))
        elif block.kind == "code":
            story.append(Preformatted(str(block.value), code, maxLineLength=116))
        elif block.kind == "image":
            alt, rel = block.value
            path = ROOT / rel
            if path.exists():
                from PIL import Image as PILImage
                with PILImage.open(path) as im:
                    w, h = im.size
                max_w = doc.width * 0.88
                max_h = 10.5 * cm
                scale = min(max_w / w, max_h / h)
                story.append(Image(str(path), width=w * scale, height=h * scale, hAlign="CENTER"))
                story.append(Paragraph(escape(alt), caption))
        elif block.kind == "table":
            rows = block.value
            if not rows:
                continue
            col_count = max(len(r) for r in rows)
            data = []
            for ri, row in enumerate(rows):
                style = table_head if ri == 0 else table_cell
                padded = row + [""] * (col_count - len(row))
                data.append([Paragraph(pdf_inline(cell), style) for cell in padded])
            table = LongTable(
                data,
                colWidths=table_widths(rows, doc.width),
                repeatRows=1,
                hAlign="CENTER",
                splitByRow=1,
            )
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#ede9fe")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#9ca3af")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#fafafa")]),
            ]))
            story.append(table)
            story.append(Spacer(1, 6))

    def decorate(canvas, document) -> None:
        canvas.saveState()
        page = canvas.getPageNumber()
        canvas.setFont("ArialReport", 7.5)
        canvas.setFillColor(colors.HexColor("#6b7280"))
        if page > 1:
            canvas.drawString(doc.leftMargin, A4[1] - 0.85 * cm, "CyberSafe — отчёт по реверсу")
        canvas.drawRightString(A4[0] - doc.rightMargin, 0.75 * cm, f"Страница {page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=decorate, onLaterPages=decorate)


def main() -> None:
    blocks = parse_markdown(SOURCE.read_text(encoding="utf-8"))
    build_docx(blocks)
    build_pdf(blocks)
    print(DOCX)
    print(PDF)


if __name__ == "__main__":
    main()
