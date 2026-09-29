#!/usr/bin/env python3
"""Offline integrity checks for the CyberSafe report package."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
import struct
import subprocess
import sys

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
EVIDENCE = ROOT / "evidence"
EXPECTED = {
    "reproduction-final/flash.bin": (2097152, "640626b93d7492a6efadba09759cf349d84a45b785cd31703ad37912d6cf9eb9"),
    "reproduction-final/pin.bin": (4, "080fe7399f197fc37d89d01ffbbb17f04f7f2ba3dd84baab2b5f09b3fc54a140"),
    "reproduction-final/disk-region.bin": (1048576, "5b07fe65994a80d2bf4cc67fd7d031ff49bfda003ae50f24a3eb41c35e57d053"),
    "reproduction-final/disk.img": (1048576, "6df635a0ddfd0f1b5c8c243382cc23dd366c63393fd18a20baaf73d5372e2c4b"),
}
MANUAL_EXPECTED = {
    "manual-button-20260927/flash.bin": (2097152, "183ec770b9581e5add2a85dc7d0fa82377bb5d92f7e525281c09250906acc4ca"),
    "manual-button-20260927/pin.bin": (4, "080fe7399f197fc37d89d01ffbbb17f04f7f2ba3dd84baab2b5f09b3fc54a140"),
    "manual-button-20260927/disk-region.bin": (1048576, "a8e537b29d87c960eac967df56e15e0474e850f6ed34fcdf92e09022b382f87f"),
    "manual-button-20260927/disk.img": (1048576, "ca31ed53d0d408e66119b5835dcfada0e40c828b974160800e44c328cd2892b6"),
}
PATCHES = [
    "bypass-compare.uf2",
    "skip-password-call.uf2",
    "pin-0000.uf2",
    "return-from-password.uf2",
]


def sha(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def check_uf2(path: Path) -> dict:
    raw = path.read_bytes()
    assert len(raw) % 512 == 0 and raw
    blocks = len(raw) // 512
    addresses = []
    for n in range(blocks):
        b = raw[n * 512:(n + 1) * 512]
        magic0, magic1, flags, address, size, number, total, family = struct.unpack_from("<8I", b)
        end, = struct.unpack_from("<I", b, 508)
        assert (magic0, magic1, end) == (0x0A324655, 0x9E5D5157, 0x0AB16F30)
        assert flags & 0x2000 and family == 0xE48BFF56
        assert size == 256 and number == n and total == blocks
        addresses.append(address)
    return {"blocks": blocks, "first": hex(min(addresses)), "last_exclusive": hex(max(addresses) + 256)}


def main() -> None:
    checks = {}
    for rel, (size, digest) in EXPECTED.items():
        p = EVIDENCE / rel
        checks[rel] = {"size": p.stat().st_size, "sha256": sha(p)}
        assert p.stat().st_size == size and sha(p) == digest

    manual_checks = {}
    for rel, (size, digest) in MANUAL_EXPECTED.items():
        p = EVIDENCE / rel
        manual_checks[rel] = {"size": p.stat().st_size, "sha256": sha(p)}
        assert p.stat().st_size == size and sha(p) == digest
    manual_dir = EVIDENCE / "manual-button-20260927"
    assert (manual_dir / "pin.bin").read_bytes() == bytes((8, 1, 7, 0))
    transition = (manual_dir / "02-button-transition.txt").read_text()
    assert "Product ID: 0x4000" in transition and "Product ID: 0x0003" in transition
    compare = json.loads((manual_dir / "compare-with-initial.json").read_text())
    assert compare["firmware_first_1MiB_identical"] and compare["pin_bytes_identical"]
    assert compare["different_bytes"] == 522 and compare["different_sector_count"] == 4
    manual_root = json.loads((manual_dir / "fat-root.json").read_text())
    assert manual_root["file_contents_read"] is False
    manual_log = (manual_dir / "session.log").read_text()
    assert "two_reads_identical\": true" in manual_log
    assert "PIN was not entered. Flash writes: none." in manual_log
    checks["manual_button_20260927"] = {
        "files": manual_checks,
        "transition": "cafe:4000 -> 2e8a:0003",
        "pin": "8170",
        "firmware_first_1MiB_identical": True,
        "fat_file_contents_read": False,
    }

    disk = (EVIDENCE / "reproduction-final/disk.img").read_bytes()
    assert disk[3:11] == b"MSDOS5.0" and disk[54:62] == b"FAT12   " and disk[510:512] == b"\x55\xaa"
    assert (EVIDENCE / "reproduction-final/pin.bin").read_bytes() == bytes((8, 1, 7, 0))
    checks["fat_boot_sector"] = {"oem": "MSDOS5.0", "filesystem": "FAT12", "signature": "55aa"}

    parsed_json = 0
    for p in EVIDENCE.rglob("*.json"):
        json.loads(p.read_text())
        parsed_json += 1
    checks["json_files_parsed"] = parsed_json

    fat_root = json.loads((EVIDENCE / "reproduction-final/fat-root.json").read_text())
    assert fat_root["file_contents_read"] is False
    ram = json.loads((EVIDENCE / "ram-exec-proof.json").read_text())
    assert all(ram[k] for k in ("pc_write_worked", "pc_exec_worked", "payload_returned_to_usb_bootloader", "original_ram_restored_and_verified"))
    transform = json.loads((EVIDENCE / "transform/transform-proof.json").read_text())
    assert transform["reencode_matches_original_ciphertext"]
    assert transform["known_plaintext_demo"]["ciphertext_xor_recovered_keystream_matches_32_known_bytes"]
    assert transform["full_sector_decode_with_firmware_transform_matches"]
    assert transform["controlled_change"]["decode_gives_requested_label"]
    assert transform["controlled_change"]["keystream_generator_used_for_modification"] is False
    checks["proof_flags"] = {"fat_contents_read": False, "ram_exec": True, "transform_roundtrip": True, "malleability": True}

    direct_decode = json.loads((EVIDENCE / "disk-from-region.img.json").read_text())
    full_dump_decode = json.loads((EVIDENCE / "disk-repeat.img.json").read_text())
    fat_metadata = json.loads((EVIDENCE / "fat-root-metadata.json").read_text())
    expected_decoded_sha = EXPECTED["reproduction-final/disk.img"][1]
    assert direct_decode["output_sha256"] == expected_decoded_sha
    assert full_dump_decode["output_sha256"] == expected_decoded_sha
    assert sha(EVIDENCE / "disk-from-region.img") == expected_decoded_sha
    assert sha(EVIDENCE / "disk-repeat.img") == expected_decoded_sha
    assert fat_metadata["filesystem_label"] == "FAT12   "
    assert fat_metadata["boot_signature_hex"] == "55aa"
    assert fat_metadata["file_contents_read"] is False
    checks["independent_decode_paths"] = {
        "direct_region_sha256": direct_decode["output_sha256"],
        "full_flash_sha256": full_dump_decode["output_sha256"],
        "identical": True,
        "fat12": True,
        "file_contents_read": False,
    }

    uf2 = {}
    for name in PATCHES:
        uf2[name] = check_uf2(EVIDENCE / "patches" / name)
        assert uf2[name]["blocks"] == 16
    checks["patch_uf2"] = uf2
    roundtrip = json.loads((EVIDENCE / "patches/roundtrip-proof.json").read_text())
    assert roundtrip["all_checks_passed"]
    assert all(item["no_other_full_image_bytes_changed"] for item in roundtrip["patches"])
    assert all(item["reconstructs_original_sector"] for item in roundtrip["restore_sectors"])
    pin_patch = next(item for item in roundtrip["patches"] if item["name"] == "pin-0000")
    assert pin_patch["hardware_tested"] and pin_patch["changed_byte_offsets"] == ["0x5500", "0x5501", "0x5502"]
    checks["patch_roundtrip"] = {"patches": 4, "restore_sectors": 2, "all_checks_passed": True}

    report = (ROOT / "CyberSafe_report.md").read_text()
    missing = [f"M{i:02d}" for i in range(1, 43) if f"M{i:02d}" not in report]
    assert not missing
    summary = json.loads((EVIDENCE / "proof-summary.json").read_text())["catalog"]
    assert len(summary["A_live_board"]) == 7
    assert len(summary["B_dump_static_or_ready_artifact"]) == 8
    assert len(summary["C_requires_more_equipment_or_build"]) == 25
    assert len(summary["N_negative"]) == 2
    categorized = (
        summary["A_live_board"]
        + summary["B_dump_static_or_ready_artifact"]
        + summary["C_requires_more_equipment_or_build"]
        + summary["N_negative"]
    )
    assert len(categorized) == len(set(categorized)) == 42
    assert set(categorized) == {f"M{i:02d}" for i in range(1, 43)}
    checks["method_catalog"] = {"variants": 42, "A": 7, "B": 8, "C": 25, "N": 2}

    live_pin = json.loads((EVIDENCE / "live-pin-0000-20260927/result.json").read_text())
    assert live_pin["live_result"]["pin_entered"] == "0000"
    assert live_pin["live_result"]["usb_vid_pid"] == "cafe:4000"
    assert live_pin["restore"]["copied_successfully"]
    no_lockout = json.loads((EVIDENCE / "live-no-lockout-20260927/result.json").read_text())
    assert no_lockout["result"]["new_usb_session_observed"]
    assert no_lockout["full_0000_to_9999_sweep_completed"] is False
    timing = json.loads((EVIDENCE / "timing-camera-20260927/result.json").read_text())
    assert timing["live_observation"]["tested"]
    assert timing["live_observation"]["prefix_timing_difference_usable"] is False
    matrix = json.loads((EVIDENCE / "proof-matrix.json").read_text())
    assert len(matrix["stages"]) == 14 and len(matrix["root_vectors"]) == 12
    for rel, item in matrix["artifact_files"].items():
        path = EVIDENCE / rel
        assert path.stat().st_size == item["size_bytes"] and sha(path) == item["sha256"]
    checks["new_live_checks"] = {
        "M20_pin_0000": True,
        "M20_restored": True,
        "M12_partial": True,
        "M11_camera_negative": True,
    }
    final_audit = json.loads((EVIDENCE / "final-audit-20260927/result.json").read_text())
    assert final_audit["current_usb"]["present"]
    assert final_audit["current_usb"]["vid_pid"] == "cafe:4000"
    assert final_audit["flash_write_performed_during_snapshot"] is False
    checks["final_device_health"] = {
        "usb_present": True,
        "vid_pid": "cafe:4000",
        "volume_mounted_in_snapshot": False,
    }
    checks["proof_matrix"] = {"stages": 14, "vectors": 12, "artifacts": len(matrix["artifact_files"])}
    assert (EVIDENCE / "PROOF_INDEX.md").is_file()
    method_matrix = json.loads((EVIDENCE / "method-proof-matrix.json").read_text())
    assert method_matrix["counts"]["methods"] == 42
    assert method_matrix["counts"]["A"] == 7
    assert method_matrix["counts"]["B"] == 8
    assert method_matrix["counts"]["C"] == 25
    assert method_matrix["counts"]["N"] == 2
    assert [item["id"] for item in method_matrix["methods"]] == [f"M{i:02d}" for i in range(1, 43)]
    assert next(item for item in method_matrix["methods"] if item["id"] == "M11")["proof_kind"] == "negative_result"
    assert next(item for item in method_matrix["methods"] if item["id"] == "M12")["proof_kind"] == "partial_live"
    for rel, item in method_matrix["artifact_index"].items():
        path = ROOT / rel
        assert path.stat().st_size == item["size_bytes"] and sha(path) == item["sha256"]
    assert (EVIDENCE / "METHOD_PROOF_MATRIX.md").is_file()
    checks["method_proof_matrix"] = method_matrix["counts"]
    assert "камера 240 FPS" in report and "организаторов" in report
    assert "M11 через камеру — N" in report
    assert "42 варианта в 12 корневых векторах" in report
    assert "матрица P01…P14" in report and "матрица P01…P13" not in report

    zip_files = [str(p.relative_to(ROOT)) for p in ROOT.rglob("*") if p.is_file() and p.suffix.lower() == ".zip"]
    assert not zip_files
    checks["zip_files_in_package"] = zip_files

    doc = Document(ROOT / "CyberSafe_LCT2026_jury_final.docx")
    assert len(doc.paragraphs) > 300 and len(doc.tables) > 5
    info = subprocess.run(["pdfinfo", str(ROOT / "CyberSafe_LCT2026_jury_final.pdf")], check=True, text=True, capture_output=True).stdout
    pages = next(int(line.split(":", 1)[1]) for line in info.splitlines() if line.startswith("Pages:"))
    assert pages >= 10
    pdf_text = subprocess.run(["pdftotext", str(ROOT / "CyberSafe_LCT2026_jury_final.pdf"), "-"], check=True, text=True, capture_output=True).stdout
    for needle in ("PIN: 8170", "PICOBOOT", "FAT12", "M42", "Содержимое ZIP не использовалось"):
        assert needle in pdf_text
    checks["documents"] = {"docx_paragraphs": len(doc.paragraphs), "docx_tables": len(doc.tables), "pdf_pages": pages}

    fsck = subprocess.run(["fsck_msdos", "-n", str(EVIDENCE / "reproduction-final/disk.img")], text=True, capture_output=True)
    assert fsck.returncode == 0
    checks["fsck_msdos"] = "passed"

    result = {"valid": True, "checks": checks}
    out = EVIDENCE / "final-validation.json"
    out.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
